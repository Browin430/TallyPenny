"""支付宝 / 微信官方账单文件的确定性解析。"""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable


class StatementParseError(ValueError):
    pass


def parse_statement_file(path: Path) -> dict[str, Any]:
    rows = _read_rows(path)
    header_index, headers, platform = _find_header(rows)
    candidates: list[dict[str, Any]] = []
    skipped = 0
    for values in rows[header_index + 1 :]:
        if not any(_clean(value) for value in values):
            continue
        row = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
        }
        try:
            candidate = (
                _parse_alipay_row(row)
                if platform == "alipay"
                else _parse_wechat_row(row)
            )
        except StatementParseError:
            candidate = None
        if candidate is None:
            skipped += 1
        else:
            candidates.append(candidate)
    if not candidates:
        raise StatementParseError("账单文件中没有可导入的收入或支出记录")
    return {
        "platform": platform,
        "candidates": candidates,
        "skipped": skipped,
    }


def _read_rows(path: Path) -> list[list[Any]]:
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text: str | None = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise StatementParseError("CSV 编码无法识别，请重新从支付宝导出")
        return [list(row) for row in csv.reader(text.splitlines())]
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise StatementParseError("服务器缺少 XLSX 解析组件") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    raise StatementParseError(f"不支持的账单文件格式: {path.suffix}")


def _find_header(rows: list[list[Any]]) -> tuple[int, list[str], str]:
    for index, row in enumerate(rows[:80]):
        headers = [_clean(value) for value in row]
        values = set(headers)
        if {"交易时间", "交易分类", "交易对方", "收/支", "金额"} <= values:
            return index, headers, "alipay"
        if {"交易时间", "交易类型", "交易对方", "收/支", "金额(元)"} <= values:
            return index, headers, "wechat"
    raise StatementParseError("没有找到支付宝或微信官方账单表头")


def _parse_alipay_row(row: dict[str, Any]) -> dict[str, Any] | None:
    category = _clean(row.get("交易分类"))
    merchant = _none_if_empty(row.get("交易对方"))
    description = _none_if_empty(row.get("商品说明"))
    direction = _clean(row.get("收/支"))
    status = _clean(row.get("交易状态"))
    combined = f"{category} {description or ''} {status}"
    if any(word in status for word in ("交易关闭", "交易失败", "已撤销")):
        return None
    is_refund = "退款" in combined
    kind = "income" if is_refund else _direction(direction)
    if kind is None:
        return None
    return _candidate(
        platform="alipay",
        kind=kind,
        amount=row.get("金额"),
        time=row.get("交易时间"),
        merchant=merchant,
        description=description,
        source_category=category,
        status=status,
        payment=row.get("收/付款方式"),
        is_refund=is_refund,
    )


def _parse_wechat_row(row: dict[str, Any]) -> dict[str, Any] | None:
    source_type = _clean(row.get("交易类型"))
    merchant = _none_if_empty(row.get("交易对方"))
    description = _none_if_empty(row.get("商品"))
    direction = _clean(row.get("收/支"))
    status = _clean(row.get("当前状态"))
    combined = f"{source_type} {description or ''} {status}"
    if any(word in status for word in ("已撤销", "支付失败", "已关闭")):
        return None
    is_refund = "退款" in combined
    kind = "income" if is_refund else _direction(direction)
    if kind is None:
        return None
    return _candidate(
        platform="wechat",
        kind=kind,
        amount=row.get("金额(元)"),
        time=row.get("交易时间"),
        merchant=merchant,
        description=description,
        source_category=source_type,
        status=status,
        payment=row.get("支付方式"),
        is_refund=is_refund,
    )


def _candidate(
    *,
    platform: str,
    kind: str,
    amount: Any,
    time: Any,
    merchant: str | None,
    description: str | None,
    source_category: str,
    status: str,
    payment: Any,
    is_refund: bool,
) -> dict[str, Any]:
    amount_cents = _amount_cents(amount)
    transaction_time = _transaction_time(time)
    text = f"{source_category} {merchant or ''} {description or ''}"
    return {
        "type": kind,
        "amountCents": amount_cents,
        "currency": "CNY",
        "categoryId": "refund" if is_refund else _category_id(text, kind),
        "subcategory": source_category[:8] or None,
        "merchant": merchant,
        "description": description,
        "transactionTime": transaction_time,
        "timeConfident": True,
        "paymentMethod": _payment_method(platform, payment),
        "note": f"{platform_label(platform)} · {status}" if status else platform_label(platform),
        "parseWarnings": [],
        "sourceType": "import",
        "confidence": 0.98,
    }


def platform_label(platform: str) -> str:
    return "支付宝账单" if platform == "alipay" else "微信支付账单"


def _direction(value: str) -> str | None:
    if value == "支出":
        return "expense"
    if value == "收入":
        return "income"
    return None


def _amount_cents(value: Any) -> int:
    cleaned = _clean(value).replace(",", "").replace("¥", "").replace("￥", "")
    try:
        cents = int((Decimal(cleaned) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError) as exc:
        raise StatementParseError(f"账单金额无法解析: {cleaned or '(空)'}") from exc
    if cents <= 0:
        raise StatementParseError("账单金额必须大于 0")
    return cents


def _transaction_time(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(value).strftime("%Y-%m-%dT%H:%M:%S")
        except (ImportError, TypeError, ValueError, OverflowError) as exc:
            raise StatementParseError(f"Excel 交易时间无法解析: {value}") from exc
    raw = _clean(value)
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, pattern).strftime("%Y-%m-%dT%H:%M:%S")
        except ValueError:
            continue
    raise StatementParseError(f"交易时间无法解析: {raw or '(空)'}")


def _payment_method(platform: str, value: Any) -> str:
    payment = _clean(value)
    if any(word in payment for word in ("银行", "信用卡", "储蓄卡")):
        return "bank_card"
    return "alipay" if platform == "alipay" else "wechat"


def _category_id(text: str, kind: str) -> str:
    if kind == "income":
        if "工资" in text:
            return "salary"
        if "红包" in text:
            return "redpacket"
        if any(word in text for word in ("理财", "收益", "利息")):
            return "invest_return"
        return "other_income"

    groups: tuple[tuple[str, Iterable[str]], ...] = (
        ("food", (
            "餐饮", "美食", "外卖", "餐厅", "餐馆", "饭店", "酒楼", "食府", "菜馆",
            "私房菜", "小吃", "快餐", "便当", "食堂", "早餐", "夜宵", "重庆小面",
            "小面", "面馆", "拉面", "牛肉面", "刀削面", "米线", "米粉", "螺蛳粉",
            "酸辣粉", "麻辣烫", "冒菜", "火锅", "串串", "烧烤", "烤肉", "烤鱼",
            "酸菜鱼", "馄饨", "云吞", "饺子", "包子", "粥铺", "盖饭", "炒饭",
            "汉堡", "炸鸡", "披萨", "寿司", "料理", "咖啡", "奶茶", "茶饮", "果汁",
            "鲜榨", "饮品", "烘焙", "蛋糕", "甜品", "面包店", "美团", "肯德基", "麦当劳",
        )),
        ("transport", ("交通", "地铁", "公交", "滴滴", "小遛", "共享电动车", "共享单车", "打车", "出行", "高铁", "机票")),
        ("shopping", ("日用百货", "数码电器", "服饰", "购物", "淘宝", "京东", "拼多多", "便利店", "螺丝钉", "螺丝刀", "螺丝", "螺母", "垫片", "紧固件", "五金", "扳手", "钳子", "钻头")),
        ("housing", ("房租", "住房", "家居", "家装", "物业")),
        ("entertainment", ("娱乐", "游戏", "电影", "文化休闲")),
        ("medical", ("医疗", "药", "医院", "健康")),
        ("education", ("教育", "培训", "书店")),
        ("travel", ("旅行", "旅游", "酒店", "住宿")),
        ("telecom", ("话费", "通讯", "流量")),
        ("utilities", ("水费", "电费", "燃气")),
        ("subscription", ("会员", "订阅", "API服务")),
        ("insurance", ("保险",)),
        ("car", ("汽车", "加油", "停车", "洗车")),
        ("pet", ("宠物",)),
        ("gift", ("礼物", "礼品")),
        ("investment", ("投资", "理财")),
    )
    for category, words in groups:
        if any(word in text for word in words):
            return category
    return "other"


def _clean(value: Any) -> str:
    return str(value if value is not None else "").replace("\t", "").strip()


def _none_if_empty(value: Any) -> str | None:
    cleaned = _clean(value)
    return None if cleaned in {"", "/"} else cleaned
